from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


WORK_DIR = Path(r"C:\robloxclone\practice_docx_work")
SOURCE = WORK_DIR / "report_original.docx"
WORK_OUTPUT = WORK_DIR / "report_fixed.docx"
FINAL_OUTPUT = Path.home() / "Downloads" / "Отчет_УТП_сетевой_модуль_Bobux_исправлено.docx"


GOST_REFERENCES = [
    "Godot Engine 4.7 Documentation [Электронный ресурс]. – URL: https://docs.godotengine.org/ (дата обращения: 25.02.2026). – Текст : электронный.",
    "PocketBase Documentation [Электронный ресурс]. – URL: https://pocketbase.io/docs/ (дата обращения: 10.04.2026). – Текст : электронный.",
    "Node.js Documentation [Электронный ресурс]. – URL: https://nodejs.org/docs/ (дата обращения: 11.03.2026). – Текст : электронный.",
    "JavaScript MDN Web Docs [Электронный ресурс]. – URL: https://developer.mozilla.org/ (дата обращения: 11.03.2026). – Текст : электронный.",
    "WebSocket API. MDN Web Docs [Электронный ресурс]. – URL: https://developer.mozilla.org/ru/docs/Web/API/WebSocket (дата обращения: 11.03.2026). – Текст : электронный.",
    "HTTP-запросы в Godot Engine [Электронный ресурс]. – URL: https://docs.godotengine.org/en/stable/classes/class_httprequest.html (дата обращения: 11.03.2026). – Текст : электронный.",
    "WebSocketPeer в Godot Engine [Электронный ресурс]. – URL: https://docs.godotengine.org/en/stable/classes/class_websocketpeer.html (дата обращения: 11.03.2026). – Текст : электронный.",
    "Express.js Documentation [Электронный ресурс]. – URL: https://expressjs.com/ (дата обращения: 11.03.2026). – Текст : электронный.",
    "PocketBase REST API [Электронный ресурс]. – URL: https://pocketbase.io/docs/api-records/ (дата обращения: 03.07.2026). – Текст : электронный.",
    "JSON Web Token Introduction [Электронный ресурс]. – URL: https://jwt.io/introduction (дата обращения: 11.03.2026). – Текст : электронный.",
    "Фаулер М. Архитектура корпоративных программных приложений. – Москва : Вильямс, 2020. – 544 с. – Текст : непосредственный.",
    "Таненбаум Э., Уэзеролл Д. Компьютерные сети. – Санкт-Петербург : Питер, 2021. – 960 с. – Текст : непосредственный.",
    "Гамма Э., Хелм Р., Джонсон Р., Влиссидес Дж. Приемы объектно-ориентированного проектирования. Паттерны проектирования. – Санкт-Петербург : Питер, 2020. – 368 с. – Текст : непосредственный.",
    "Ричардсон К. Микросервисы. Паттерны разработки и рефакторинга. – Санкт-Петербург : Питер, 2019. – 544 с. – Текст : непосредственный.",
    "Нильсен Я. Веб-дизайн: удобство использования Web-сайтов. – Москва : Вильямс, 2020. – 368 с. – Текст : непосредственный.",
    "Маклафлин Б. Объектно-ориентированный анализ и проектирование. – Санкт-Петербург : Питер, 2021. – 624 с. – Текст : непосредственный.",
    "Субботин С. А. Сетевые технологии и распределенные информационные системы : учебное пособие. – Москва : Инфра-М, 2022. – 312 с. – Текст : непосредственный.",
    "Гроховский В. Л. Базы данных и информационные системы : учебное пособие. – Москва : Юрайт, 2023. – 256 с. – Текст : непосредственный.",
    "REST API Tutorial [Электронный ресурс]. – URL: https://restfulapi.net/ (дата обращения: 03.07.2026). – Текст : электронный.",
    "SQLite Documentation [Электронный ресурс]. – URL: https://sqlite.org/docs.html (дата обращения: 03.07.2026). – Текст : электронный.",
    "OWASP Top Ten [Электронный ресурс]. – URL: https://owasp.org/www-project-top-ten/ (дата обращения: 03.07.2026). – Текст : электронный.",
]


def set_times_new_roman(paragraph, size_pt=14, bold=None) -> None:
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(size_pt)
        if bold is not None:
            run.font.bold = bold


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    shutil.copy2(SOURCE, WORK_OUTPUT)
    doc = Document(WORK_OUTPUT)

    section_one_index = None
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "1. Описание постановки задачи":
            section_one_index = index
            paragraph.paragraph_format.page_break_before = True
            break
    else:
        raise RuntimeError("Не найден раздел 1")

    # LibreOffice may ignore page_break_before for imported Normal paragraphs,
    # so add an explicit page break at the end of the previous paragraph too.
    if section_one_index and section_one_index > 0:
        doc.paragraphs[section_one_index - 1].add_run().add_break(WD_BREAK.PAGE)

    bibliography_index = None
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "Список литературы":
            bibliography_index = index
            break
    if bibliography_index is None:
        raise RuntimeError("Не найден список литературы")

    for paragraph in doc.paragraphs[bibliography_index + 1:]:
        element = paragraph._element
        element.getparent().remove(element)

    heading = doc.paragraphs[bibliography_index]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_times_new_roman(heading, bold=True)

    for number, reference in enumerate(GOST_REFERENCES, 1):
        paragraph = doc.add_paragraph(f"{number}. {reference}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = paragraph.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.left_indent = Cm(0)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.5
        set_times_new_roman(paragraph)

    # The two screenshots that form Figure 7 were wider than the printable area.
    for shape_index in (6, 7):
        shape = doc.inline_shapes[shape_index]
        old_width = shape.width
        old_height = shape.height
        new_width = Cm(15.0)
        shape.width = new_width
        shape.height = int(old_height * (new_width / old_width))

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(WORK_OUTPUT)
    shutil.copy2(WORK_OUTPUT, FINAL_OUTPUT)
    print(f"SAVED_WORK={WORK_OUTPUT}")
    print(f"SAVED_FINAL={FINAL_OUTPUT}")


if __name__ == "__main__":
    main()
