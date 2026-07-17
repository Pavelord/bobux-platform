from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path("C:/robloxclone")
PRACTICE_DIR = ROOT / "практика"
SOURCE_REPORT = PRACTICE_DIR / "Отчет_УТП_сетевой_модуль_Bobux_по_образцу.docx"
OUT_DOCX = PRACTICE_DIR / "Отчет_УТП_сетевой_модуль_Bobux_точная_копия_примера.docx"


SECTION_TITLES = [
    "Введение",
    "1. Описание постановки задачи",
    "2. Структура ролей пользователей",
    "3. Обоснование выбора технологий",
    "4. Описание архитектуры сетевого модуля",
    "5. Описание программных модулей регистрации и авторизации",
    "6. Описание программного модуля пользовательского контента",
    "7. Описание программного модуля игровых серверов",
    "8. Описание административного модуля статистики",
    "9. Тестирование сетевого модуля",
    "Заключение",
    "Список литературы",
    "Приложение А. Основные файлы проекта",
    "Приложение Б. Сценарий демонстрации",
    "Приложение В. План иллюстраций и скриншотов",
]

TOC_ENTRIES = [
    ("Введение", "3"),
    ("1. Описание постановки задачи", "5"),
    ("2. Структура ролей пользователей", "7"),
    ("3. Обоснование выбора технологий", "8"),
    ("4. Описание архитектуры сетевого модуля", "9"),
    ("5. Описание программных модулей регистрации и авторизации", "10"),
    ("6. Описание программного модуля пользовательского контента", "11"),
    ("7. Описание программного модуля игровых серверов", "13"),
    ("8. Описание административного модуля статистики", "14"),
    ("9. Тестирование сетевого модуля", "15"),
    ("Заключение", "16"),
    ("Список литературы", "17"),
    ("Приложение А. Основные файлы проекта", "18"),
    ("Приложение Б. Сценарий демонстрации", "19"),
    ("Приложение В. План иллюстраций и скриншотов", "20"),
]

FIGURE_CAPTIONS = {
    "5. Описание программных модулей регистрации и авторизации": [
        "Рисунок 1 - Форма авторизации пользователя в клиенте Bobux.",
        "Рисунок 2 - Серверные маршруты регистрации и авторизации.",
        "Рисунок 3 - Клиентский сценарий авторизации и синхронизации профиля.",
    ],
    "6. Описание программного модуля пользовательского контента": [
        "Рисунок 4 - Отображение пользовательского каталога и инвентаря в клиенте.",
        "Рисунок 5 - Серверная обработка публикации предметов каталога.",
    ],
    "7. Описание программного модуля игровых серверов": [
        "Рисунок 6 - WebSocket-подключение к выделенному игровому серверу.",
        "Рисунок 7 - Heartbeat активных игровых комнат в базе данных.",
    ],
    "8. Описание административного модуля статистики": [
        "Рисунок 8 - Статистика и коллекции базы данных сетевого модуля.",
        "Рисунок 9 - Формирование административной статистики проекта Bobux.",
    ],
    "9. Тестирование сетевого модуля": [
        "Рисунок 10 - Главное меню платформы после загрузки профиля пользователя.",
    ],
}


def set_run_font(run, size: float = 12, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color or RGBColor(0, 0, 0)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    if "toc 1" in doc.styles:
        toc = doc.styles["toc 1"]
        toc.font.name = "Times New Roman"
        toc._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        toc._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        toc.font.size = Pt(12)
        toc.paragraph_format.space_before = Pt(0)
        toc.paragraph_format.space_after = Pt(0)


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    align: WD_ALIGN_PARAGRAPH | None = None,
    first_indent: bool = True,
    size: float = 12,
    bold: bool = False,
    line_spacing: float = 1.5,
    color: RGBColor | None = None,
):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Cm(1.25 if first_indent else 0)
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_blank(doc: Document, count: int = 1) -> None:
    for _ in range(count):
        add_paragraph(doc, "", first_indent=False)


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_bottom_border(paragraph, size: str = "18") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)


def add_title_page(doc: Document) -> None:
    add_blank(doc, 1)
    add_paragraph(
        doc,
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯРОССИЙСКОЙ\nФЕДЕРАЦИИ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        bold=True,
        line_spacing=1.0,
    )
    add_paragraph(
        doc,
        "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ\nУЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        bold=True,
        line_spacing=1.0,
    )
    add_paragraph(
        doc,
        "«БАЛТИЙСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ ИМЕНИ ИММАНУИЛА\nКАНТА»",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
        bold=True,
        line_spacing=1.0,
    )
    line = add_paragraph(doc, "", first_indent=False)
    add_bottom_border(line)
    add_blank(doc, 7)
    add_paragraph(doc, "ОТЧЕТ", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, size=14)
    add_paragraph(doc, "по учебно-технологической практике", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    add_paragraph(doc, "студента 1 курса группы [номер группы]", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    add_blank(doc, 1)
    add_paragraph(
        doc,
        "специальности 01.03.02 Прикладная математика и информатика",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_indent=False,
    )
    add_paragraph(doc, "[ФИО студента]", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, color=RGBColor(192, 0, 0))
    add_blank(doc, 5)
    add_paragraph(doc, "Отметка о защите отчета", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, line_spacing=1.0)
    add_blank(doc, 1)
    add_paragraph(doc, "Отчет защищен с оценкой ____________________", first_indent=False, line_spacing=1.0)
    add_paragraph(doc, "«___»____________ 20____г.", first_indent=False, line_spacing=1.0)
    add_blank(doc, 3)
    add_paragraph(
        doc,
        "Руководитель учебно-\nтехнологической практики    ________________________      [ФИО руководителя]",
        first_indent=False,
        line_spacing=1.0,
    )
    add_blank(doc, 1)
    add_paragraph(doc, "Калининград 2026", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
    add_page_break(doc)


def add_toc(doc: Document) -> None:
    add_blank(doc, 2)
    add_paragraph(doc, "Содержание", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
    add_blank(doc, 2)
    for title, page in TOC_ENTRIES:
        paragraph = doc.add_paragraph(style="toc 1" if "toc 1" in doc.styles else None)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = paragraph.add_run(f"{title}\t{page}")
        set_run_font(run, size=12)
    add_page_break(doc)


def extract_source_paragraphs() -> list[str]:
    source = Document(SOURCE_REPORT)
    result: list[str] = []
    started = False
    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if text == "Введение":
            started = True
        if not started:
            continue
        if text in ("Содержание",):
            continue
        if "\t" in text:
            continue
        result.append(text)
    return result


def add_body_from_source(doc: Document) -> None:
    paragraphs = extract_source_paragraphs()
    current_heading = ""
    inserted_caption_for_heading: set[str] = set()
    paragraph_count_under_heading = 0
    for text in paragraphs:
        is_heading = text in SECTION_TITLES
        if is_heading:
            current_heading = text
            paragraph_count_under_heading = 0
            if text not in ("Введение",):
                add_page_break(doc)
            add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, bold=True)
            continue

        if text.startswith("Рисунок ") or text.startswith("[МЕСТО"):
            continue

        align = WD_ALIGN_PARAGRAPH.JUSTIFY
        first_indent = True
        if current_heading == "Список литературы":
            first_indent = False
        add_paragraph(doc, text, align=align, first_indent=first_indent)
        paragraph_count_under_heading += 1

        if current_heading in FIGURE_CAPTIONS and current_heading not in inserted_caption_for_heading and paragraph_count_under_heading == 2:
            inserted_caption_for_heading.add(current_heading)
            for caption in FIGURE_CAPTIONS[current_heading]:
                add_blank(doc, 2)
                add_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, line_spacing=1.0)


def add_extra_screenshot_plan(doc: Document) -> None:
    add_blank(doc, 1)
    add_paragraph(
        doc,
        "Для усиления доказательной базы отчета рекомендуется вставить скриншоты не только интерфейса, но и ключевых фрагментов исходного кода. Ниже приведен расширенный перечень иллюстраций. Если рисунок вставляется вручную, его следует размещать непосредственно над соответствующей подписью.",
    )
    items = [
        ("Экран входа в приложение Bobux до авторизации пользователя.", "Рисунок 11 - Экран входа пользователя в клиент Bobux."),
        ("Главное меню после успешной авторизации в учетной записи pavelord или тестовом аккаунте.", "Рисунок 12 - Загрузка профиля и основных разделов после входа в аккаунт."),
        ("Фрагмент файла C:/robloxclone/services/bobux_api/server.js со строками регистрации и входа.", "Рисунок 13 - Серверные маршруты регистрации и авторизации пользователя."),
        ("Фрагмент файла C:/robloxclone/autoload/cloud_api.gd с функциями sign_up_with_credentials и sign_in_with_credentials.", "Рисунок 14 - Клиентские функции регистрации и входа в сетевой модуль."),
        ("Фрагмент файла C:/robloxclone/autoload/cloud_api.gd с универсальной функцией _request_http_json.", "Рисунок 15 - Универсальный HTTP-транспорт клиентского сетевого модуля."),
        ("Фрагмент файла C:/robloxclone/services/bobux_api/server.js с publish_avatar_item и fetch_my_creations.", "Рисунок 16 - Серверная обработка пользовательского каталога и созданных предметов."),
        ("Фрагмент файла C:/robloxclone/autoload/client_network.gd с WebSocket-подключением и подтверждением входа в комнату.", "Рисунок 17 - Подключение клиента к выделенной игровой комнате."),
        ("Фрагмент файла C:/robloxclone/server/server_main.gd с heartbeat активного сервера.", "Рисунок 18 - Регистрация активной игровой комнаты на сервере."),
        ("Окно PocketBase или административная статистика с количеством пользователей, карт, предметов и записей инвентаря.", "Рисунок 19 - Фактическое состояние базы данных проекта Bobux."),
        ("Раздел каталога или инвентаря, где отображается пользовательский контент, загруженный с сервера.", "Рисунок 20 - Отображение серверного пользовательского контента в клиентском интерфейсе."),
    ]
    for description, caption in items:
        add_paragraph(doc, description, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True)
        add_blank(doc, 2)
        add_paragraph(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, line_spacing=1.0)
        add_blank(doc, 1)


def build() -> None:
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_toc(doc)
    add_body_from_source(doc)
    add_extra_screenshot_plan(doc)
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build()
