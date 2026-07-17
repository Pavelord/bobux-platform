from __future__ import annotations

import json
import math
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\robloxclone")
PRACTICE_DIR = next(
    path for path in ROOT.iterdir()
    if path.is_dir() and any(ord(char) > 127 for char in path.name)
)
STATS_PATH = PRACTICE_DIR / "bobux_network_module_stats_2026-07-03.json"
REPORT_DOCX = PRACTICE_DIR / "Отчет_УТП_сетевой_модуль_Bobux.docx"
REPORT_MD = PRACTICE_DIR / "Отчет_УТП_сетевой_модуль_Bobux.md"
GUIDE_MD = PRACTICE_DIR / "Памятка_защита_сетевой_модуль_Bobux.md"
DIAGRAM_PATH = PRACTICE_DIR / "architecture_network_module.png"


def load_stats() -> dict:
    return json.loads(STATS_PATH.read_text(encoding="utf-8"))


def stat_summary(stats: dict) -> dict:
    return {
        "auth_users": stats["accounts"]["auth_users"],
        "profiles": stats["accounts"]["profiles"],
        "named_profiles": stats["accounts"]["named_profiles"],
        "maps": stats["creation"]["maps"],
        "map_creators": stats["creation"]["map_creators"],
        "total_visits": stats["creation"]["total_visits"],
        "total_likes": stats["creation"]["total_likes"],
        "model_assets": stats["creation"]["model_assets"],
        "avatar_items": stats["creation"]["avatar_items"],
        "inventory": stats["creation"]["inventory_records"],
        "friendships": stats["social"]["friendships"],
        "pb_mb": stats["database"]["pocketbase_data_mb"],
        "storage_mb": stats["storage"]["public_storage_mb"],
    }


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def build_architecture_diagram() -> None:
    width, height = 1600, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_title = load_font(34, True)
    font_h = load_font(23, True)
    font = load_font(19)
    font_small = load_font(16)
    colors = {
        "blue": "#2E74B5",
        "dark": "#1F3A5F",
        "light": "#E8EEF5",
        "green": "#E8F5E9",
        "orange": "#FFF3E0",
        "gray": "#F4F6F9",
        "border": "#8EA9C1",
    }

    def box(x1: int, y1: int, x2: int, y2: int, title: str, lines: list[str], fill: str) -> None:
        draw.rounded_rectangle([x1, y1, x2, y2], radius=22, fill=fill, outline=colors["border"], width=3)
        draw.text((x1 + 24, y1 + 18), title, fill=colors["dark"], font=font_h)
        y = y1 + 58
        for line in lines:
            draw.text((x1 + 28, y), line, fill="#1F1F1F", font=font)
            y += 31

    def arrow(x1: int, y1: int, x2: int, y2: int, label: str = "") -> None:
        draw.line([x1, y1, x2, y2], fill=colors["blue"], width=4)
        angle = math.atan2(y2 - y1, x2 - x1)
        length = 18
        for arrow_angle in [angle + 2.6, angle - 2.6]:
            draw.line(
                [x2, y2, x2 - length * math.cos(arrow_angle), y2 - length * math.sin(arrow_angle)],
                fill=colors["blue"],
                width=4,
            )
        if label:
            mx, my = (x1 + x2) // 2, (y1 + y2) // 2
            draw.rounded_rectangle([mx - 90, my - 20, mx + 90, my + 18], radius=10, fill="white", outline="#D0D7DE")
            draw.text((mx - 78, my - 14), label, fill=colors["dark"], font=font_small)

    draw.text((80, 40), "Архитектура сетевого модуля игровой платформы Bobux", fill=colors["dark"], font=font_title)
    box(70, 140, 380, 335, "Клиент Godot", ["Login / Lobby", "CloudAPI.gd", "NetworkManager.gd", "Avatar / Catalog UI"], colors["light"])
    box(70, 505, 380, 700, "Лаунчер", ["Проверка версии", "Загрузка сборки", "Манифесты latest.json", "Windows / Android"], colors["gray"])
    box(520, 140, 860, 335, "HTTP API Node.js", ["Авторизация", "CRUD данных", "Валидация", "Админ-статистика"], colors["green"])
    box(520, 505, 860, 700, "WebSocket-сервер", ["Игровые комнаты", "Сессии игроков", "Передача событий", "Heartbeat"], colors["orange"])
    box(1010, 120, 1460, 300, "PocketBase", ["profiles, maps, avatar_items", "friendships, active_servers", "avatar_outfits, inventory", "хранение бизнес-данных"], colors["light"])
    box(1010, 360, 1460, 540, "Файловое хранилище", ["превью карт", "модели и ассеты", "текстуры одежды", "публичные URL"], colors["gray"])
    box(1010, 610, 1460, 790, "Nginx + PM2 + VPS", ["маршрутизация /api и /ws", "перезапуск сервисов", "публичный доступ", "единая точка входа"], "#F7F7F7")
    arrow(380, 220, 520, 220, "REST")
    arrow(380, 610, 520, 610, "manifest")
    arrow(690, 335, 690, 505, "WS")
    arrow(860, 220, 1010, 210, "PB API")
    arrow(860, 265, 1010, 440, "storage")
    arrow(860, 610, 1010, 700, "proxy")
    arrow(1240, 300, 1240, 360, "files")
    image.save(DIAGRAM_PATH)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True)
        set_cell_shading(table.rows[0].cells[index], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    return table


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_para(doc: Document, text: str) -> None:
    for part in textwrap.dedent(text).strip().split("\n\n"):
        doc.add_paragraph(part.strip())


def configure_doc_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.1
    for name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)


def write_report_markdown(stats_lines: dict) -> None:
    text = f"""# Отчёт по учебно-технологической практике

**Тема:** Разработка сетевого модуля игровой платформы Bobux с авторизацией, каталогом пользовательского контента и синхронизацией игровых данных

**Студент:** [ФИО студента]

**Группа:** [номер группы]

**Руководитель:** [ФИО руководителя]

**Год:** 2026

## Аннотация

В рамках практики разработан сетевой модуль игровой платформы Bobux. Модуль отвечает за авторизацию пользователей, хранение профилей, каталог пользовательского контента, работу с друзьями, синхронизацию аватара, регистрацию активных игровых серверов, хранение файлов и административную статистику.

## Фактические показатели системы на 03.07.2026

- auth-пользователи: {stats_lines['auth_users']}
- профили игроков: {stats_lines['profiles']}
- пользовательские карты: {stats_lines['maps']}
- создатели карт: {stats_lines['map_creators']}
- посещения карт: {stats_lines['total_visits']}
- лайки карт: {stats_lines['total_likes']}
- модели: {stats_lines['model_assets']}
- предметы аватара: {stats_lines['avatar_items']}
- записи инвентаря: {stats_lines['inventory']}
- дружеские связи: {stats_lines['friendships']}
- размер PocketBase data: {stats_lines['pb_mb']} MB
- размер публичного storage: {stats_lines['storage_mb']} MB

Полная версия отчёта сохранена в DOCX-файле.
"""
    REPORT_MD.write_text(text, encoding="utf-8")


def add_title_page(doc: Document) -> None:
    for line in [
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ",
        "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
        "«БАЛТИЙСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ ИМЕНИ ИММАНУИЛА КАНТА»",
    ]:
        paragraph = doc.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.runs[0].font.size = Pt(10)
        paragraph.runs[0].bold = True
    doc.add_paragraph("\n\n")
    paragraph = doc.add_paragraph("ОТЧЁТ")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].font.size = Pt(18)
    paragraph.runs[0].bold = True
    paragraph = doc.add_paragraph("по учебно-технологической практике")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].font.size = Pt(14)
    doc.add_paragraph("\n")
    paragraph = doc.add_paragraph(
        "Тема: «Разработка сетевого модуля игровой платформы Bobux "
        "с авторизацией, каталогом пользовательского контента и синхронизацией игровых данных»"
    )
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.runs[0].font.size = Pt(13)
    paragraph.runs[0].bold = True
    for line in ["\nСтудент: [ФИО студента]", "Группа: [номер группы]", "Руководитель практики: [ФИО руководителя]"]:
        paragraph = doc.add_paragraph(line)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph("\n\n\n")
    paragraph = doc.add_paragraph("Калининград 2026")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_paragraph("Содержание", style="Heading 1")
    for item in [
        "Введение",
        "1. Описание постановки задачи",
        "2. Структура ролей пользователей",
        "3. Обоснование выбора технологий",
        "4. Общая архитектура сетевого модуля",
        "5. База данных и основные коллекции",
        "6. Реализация авторизации и профилей",
        "7. Каталог пользовательского контента",
        "8. Работа с файлами и изображениями",
        "9. Социальный модуль",
        "10. Регистрация активных игровых серверов",
        "11. Сетевое взаимодействие клиента и сервера",
        "12. Серверная валидация и безопасность",
        "13. Административная статистика",
        "14. Тестирование",
        "15. Результаты работы",
        "Заключение",
        "Список использованных источников",
        "Приложения",
    ]:
        doc.add_paragraph(item)
    doc.add_page_break()


def build_report_docx(stats_lines: dict) -> None:
    doc = Document()
    configure_doc_styles(doc)
    add_title_page(doc)
    add_contents(doc)

    doc.add_paragraph("Введение", style="Heading 1")
    add_para(doc, """
    Целью учебно-технологической практики является получение практического опыта разработки программных модулей, работающих с базой данных, пользовательскими аккаунтами, файлами и сетевыми запросами. В качестве предметной области выбрана игровая платформа Bobux. Платформа включает клиент на Godot, серверный API, базу данных PocketBase, файловое хранилище, выделенный игровой сервер и лаунчер обновлений.

    В рамках практики рассматривается не вся игровая платформа целиком, а отдельный законченный программный модуль: сетевой модуль игровой платформы. Он отвечает за регистрацию и авторизацию пользователей, хранение профилей, каталог пользовательских игр и предметов, синхронизацию аватара, работу с друзьями, регистрацию активных игровых серверов и загрузку файлов.

    При разработке применялись нейросетевые агенты как вспомогательный инструмент для анализа кода, поиска ошибок и ускорения реализации. Постановка задачи, выбор архитектуры, интеграция, проверка и итоговое оформление выполнялись автором проекта.
    """)

    doc.add_paragraph("1. Описание постановки задачи", style="Heading 1")
    add_para(doc, """
    Необходимо разработать сетевой модуль для игровой платформы, который обеспечивает взаимодействие игрового клиента с серверной инфраструктурой. Модуль должен выполнять регистрацию и авторизацию пользователей, хранить профили и данные аватара, обеспечивать CRUD-операции для пользовательского контента, хранить карты, модели и предметы аватара, поддерживать список друзей, регистрировать активные игровые серверы и предоставлять административную статистику.

    Поставленная задача соответствует общим требованиям учебно-технологической практики: используются база данных, авторизация, управление записями, загрузка файлов, поиск и серверная проверка данных.
    """)

    doc.add_paragraph("2. Структура ролей пользователей", style="Heading 1")
    add_table(doc, ["Роль", "Возможности"], [
        ["Неавторизованный пользователь", "Открытие клиента, переход к регистрации или авторизации"],
        ["Авторизованный игрок", "Лобби, запуск игр, профиль, аватар, инвентарь, друзья"],
        ["Создатель контента", "Публикация карт, моделей и предметов аватара"],
        ["Администратор", "Доступ к защищённой статистике и состоянию сервера"],
    ])

    doc.add_paragraph("3. Обоснование выбора технологий", style="Heading 1")
    add_table(doc, ["Технология", "Назначение"], [
        ["Godot 4.7 / GDScript", "Клиент игровой платформы, UI, лобби и подключение к игровому серверу"],
        ["Node.js / Express", "Серверный HTTP API, валидация, маршрутизация запросов"],
        ["PocketBase", "База данных и коллекции пользовательских данных"],
        ["WebSocket", "Сетевое взаимодействие в игровой сессии"],
        ["Nginx", "Проксирование HTTP и WebSocket-запросов"],
        ["PM2", "Запуск и восстановление Node.js API на VPS"],
    ])

    doc.add_paragraph("4. Общая архитектура сетевого модуля", style="Heading 1")
    add_para(doc, """
    Сетевой модуль построен по многоуровневой архитектуре. Клиент Godot не обращается к базе данных напрямую. Все операции выполняются через серверный API. Это повышает безопасность и позволяет централизованно контролировать валидацию, права доступа, формат данных и загрузку файлов.

    Основные уровни архитектуры: клиентский уровень Godot, серверный API Node.js, база данных PocketBase, файловое хранилище, игровой WebSocket-сервер и инфраструктурный слой VPS.
    """)
    doc.add_picture(str(DIAGRAM_PATH), width=Inches(6.4))
    caption = doc.add_paragraph("Рисунок 1 — Архитектура сетевого модуля игровой платформы Bobux")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("5. База данных и основные коллекции", style="Heading 1")
    add_table(doc, ["Коллекция", "Назначение"], [
        ["profiles", "Профили игроков, имя, статус, текущая игра, данные аватара"],
        ["maps", "Пользовательские карты и игровые режимы"],
        ["active_servers", "Активные игровые комнаты и heartbeat-состояние серверов"],
        ["model_assets", "Загруженные пользовательские модели"],
        ["avatar_items", "Одежда, аксессуары и предметы аватара"],
        ["avatar_outfits", "Текущее состояние внешнего вида пользователя"],
        ["user_inventory", "Связь пользователя с полученными предметами"],
        ["friendships", "Подтверждённые дружеские связи"],
        ["friend_requests", "Входящие и исходящие заявки в друзья"],
        ["follows", "Подписки между пользователями"],
        ["asset_likes", "Лайки карт, моделей и предметов"],
    ])

    doc.add_paragraph("6. Реализация авторизации и профилей", style="Heading 1")
    add_para(doc, """
    Авторизация реализована через совместимый API-слой. Клиент отправляет логин и пароль на серверный маршрут /api/auth/v1/token. Сервер проверяет пользователя через PocketBase и возвращает access token. Клиент сохраняет сессию локально и использует bearer token для последующих запросов.

    После входа выполняется загрузка профиля пользователя. В профиле хранится имя, статус, текущая игра и базовые данные аватара. Отдельно хранится avatar_outfits, где содержится актуальная одежда, цвета тела и список надетых предметов. Такое разделение позволяет не перегружать основной профиль тяжёлыми данными и быстрее открывать лобби.
    """)
    add_bullets(doc, [
        "загрузка outfit после входа выполняется фоном и не блокирует логин;",
        "публичные списки друзей получают облегчённые профили без тяжёлых inline-данных;",
        "сервер обрезает слишком крупные thumbnails в ответах списка карт и профилей;",
        "общий HTTP timeout уменьшен, чтобы зависший запрос не держал интерфейс слишком долго.",
    ])

    doc.add_paragraph("7. Каталог пользовательского контента", style="Heading 1")
    add_para(doc, """
    Модуль поддерживает карты и игровые режимы, модели, предметы аватара, изображения и превью, элементы инвентаря. Публикация контента проходит через серверный API. Клиент отправляет данные, сервер проверяет пользователя, сохраняет запись в нужной коллекции и при необходимости сохраняет файл в публичное хранилище.
    """)

    doc.add_paragraph("8. Работа с файлами и изображениями", style="Heading 1")
    add_para(doc, """
    Файлы не хранятся напрямую внутри базы данных в виде больших бинарных объектов. Для них используется файловое хранилище, а в базе сохраняются пути или публичные URL. Такой подход снижает нагрузку на базу данных и ускоряет выдачу списков. Сервер дополнительно контролирует размер inline-данных, чтобы старые или слишком крупные изображения не замораживали клиентский интерфейс.
    """)

    doc.add_paragraph("9. Социальный модуль", style="Heading 1")
    add_para(doc, """
    Социальная часть сетевого модуля включает отправку заявки в друзья, получение входящих и исходящих заявок, принятие и отклонение заявки, удаление друга, список друзей и подсчёт социальных связей. Социальные RPC-маршруты работают через серверный access token. Клиент не получает доступа к базе напрямую.
    """)

    doc.add_paragraph("10. Регистрация активных игровых серверов", style="Heading 1")
    add_para(doc, """
    Для отображения активных игровых режимов используется коллекция active_servers. Игровой сервер или клиентский хост отправляет heartbeat-записи, где указывается карта, комната, количество игроков, хост и время последней активности. Лобби запрашивает список активных серверов и может показывать, где прямо сейчас есть игроки.
    """)

    doc.add_paragraph("11. Сетевое взаимодействие клиента и сервера", style="Heading 1")
    add_para(doc, """
    В клиенте Godot сетевой модуль разделён на две основные части: CloudAPI.gd выполняет HTTP-операции, а NetworkManager отвечает за WebSocket-подключение к игровой сессии. HTTP используется для долговременных данных, WebSocket — для событий игрового процесса в реальном времени.
    """)
    add_numbered(doc, [
        "Пользователь вводит логин и пароль.",
        "Клиент отправляет запрос на /api/auth/v1/token.",
        "Сервер проверяет данные через PocketBase.",
        "Клиент получает токен и загружает профиль.",
        "Лобби открывается без ожидания тяжёлых данных.",
        "Одежда и дополнительные данные догружаются в фоне.",
        "При запуске режима клиент подключается к WebSocket-серверу.",
    ])

    doc.add_paragraph("12. Серверная валидация и безопасность", style="Heading 1")
    add_bullets(doc, [
        "клиент не имеет прямого доступа к PocketBase;",
        "все операции проходят через Express API;",
        "административная статистика защищена superuser-токеном;",
        "service credentials не передаются клиенту;",
        "загрузка файлов проходит через серверные маршруты;",
        "пути файлов нормализуются и сохраняются в безопасной директории;",
        "тяжёлые inline-данные очищаются перед выдачей публичным спискам;",
        "сервер проверяет текущего пользователя по bearer token.",
    ])

    doc.add_paragraph("13. Административная статистика и фактическое состояние системы", style="Heading 1")
    add_para(doc, """
    Для проверки состояния сетевого модуля реализован защищённый endpoint /api/admin/stats. Он собирает статистику по аккаунтам, профилям, картам, предметам, инвентарю, дружбам, активным серверам и размеру хранилищ.
    """)
    add_table(doc, ["Показатель", "Значение"], [
        ["Зарегистрированные auth-пользователи", stats_lines["auth_users"]],
        ["Профили игроков", stats_lines["profiles"]],
        ["Именованные профили", stats_lines["named_profiles"]],
        ["Пользовательские карты", stats_lines["maps"]],
        ["Создатели карт", stats_lines["map_creators"]],
        ["Суммарные посещения карт", stats_lines["total_visits"]],
        ["Суммарные лайки карт", stats_lines["total_likes"]],
        ["Загруженные модели", stats_lines["model_assets"]],
        ["Предметы аватара", stats_lines["avatar_items"]],
        ["Записи инвентаря", stats_lines["inventory"]],
        ["Дружеские связи", stats_lines["friendships"]],
        ["Размер данных PocketBase", f"{stats_lines['pb_mb']} MB"],
        ["Размер публичного файлового хранилища", f"{stats_lines['storage_mb']} MB"],
    ])

    doc.add_paragraph("14. Тестирование", style="Heading 1")
    add_bullets(doc, [
        "синтаксическая проверка Node.js API командой node --check server.js;",
        "проверка проекта Godot через Godot --headless --check-only;",
        "проверка доступности /api/health;",
        "проверка защищённой статистики /api/admin/stats;",
        "проверка загрузки списков профилей и карт без тяжёлых inline-данных;",
        "ручная проверка входа, лобби, каталога, друзей и аватара.",
    ])

    doc.add_paragraph("15. Результаты работы", style="Heading 1")
    add_bullets(doc, [
        "реализованы авторизация и управление сессиями;",
        "организовано хранение пользовательских профилей;",
        "создан каталог карт, моделей и предметов;",
        "реализована загрузка и выдача файлов;",
        "добавлена синхронизация аватара и инвентаря;",
        "реализована социальная система друзей;",
        "добавлена регистрация активных игровых серверов;",
        "добавлена защищённая административная статистика;",
        "оптимизированы тяжёлые сетевые ответы.",
    ])

    doc.add_paragraph("Заключение", style="Heading 1")
    add_para(doc, """
    Разработанный сетевой модуль решает основные задачи серверной части игровой платформы: хранение данных, авторизация, синхронизация клиента и сервера, публикация пользовательского контента и администрирование. Работа позволила применить навыки проектирования базы данных, разработки API, работы с сетевыми протоколами, файловым хранилищем и клиент-серверной архитектурой.

    В дальнейшем модуль можно развивать: добавить полноценную веб-админку, расширить систему ролей, добавить модерацию пользовательского контента, улучшить мониторинг активных игровых серверов и внедрить автоматическое резервное копирование.
    """)

    doc.add_paragraph("Список использованных источников", style="Heading 1")
    for source in [
        "Методические материалы по учебно-технологической практике, предоставленные руководителем практики.",
        "Документация Godot Engine: https://docs.godotengine.org/",
        "Документация Node.js: https://nodejs.org/docs/",
        "Документация Express: https://expressjs.com/",
        "Документация PocketBase: https://pocketbase.io/docs/",
        "MDN Web Docs: материалы по HTTP, REST и WebSocket.",
    ]:
        doc.add_paragraph(source, style="List Number")

    doc.add_paragraph("Приложение А. Основные файлы проекта", style="Heading 1")
    add_table(doc, ["Файл", "Назначение"], [
        ["autoload/cloud_api.gd", "HTTP-клиент Godot для работы с серверным API"],
        ["autoload/client_network.gd", "WebSocket-клиент и подключение к игровой сессии"],
        ["autoload/user_session.gd", "Локальная пользовательская сессия"],
        ["autoload/game_state.gd", "Состояние аватара и выбранной карты"],
        ["services/bobux_api/server.js", "Основной Node.js/Express API"],
        ["server/server_main.gd", "Выделенный игровой WebSocket-сервер"],
        ["scripts/lobby/lobby.gd", "Лобби, каталог, профиль и социальные разделы"],
        ["tools/deploy_bobux_release.ps1", "Скрипт подготовки релиза"],
        ["ops/vps/bobux_nginx.conf", "Конфигурация Nginx для API и WebSocket"],
    ])

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = "Отчёт по учебно-технологической практике — сетевой модуль Bobux"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(9)
        footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    doc.save(REPORT_DOCX)


def write_defense_guide(stats_lines: dict) -> None:
    text = f"""# Памятка для защиты: сетевой модуль Bobux

## Короткая формулировка проекта

Я разработал сетевой модуль для игровой платформы Bobux. Этот модуль отвечает за авторизацию пользователей, хранение профилей, каталог пользовательского контента, синхронизацию аватаров, работу друзей, регистрацию активных игровых серверов и загрузку файлов.

Важно говорить именно **модуль**, а не «я полностью сделал Roblox». Игровая платформа — большая среда, а твоя защищаемая часть — серверно-сетевой слой.

## Как объяснить за 30 секунд

Bobux — это игровая платформа с клиентом на Godot. Чтобы игроки могли входить в аккаунт, сохранять аватар, публиковать карты, видеть каталог игр и подключаться к серверам, нужен сетевой модуль. Я реализовал связку Godot-клиент -> Node.js API -> PocketBase -> файловое хранилище -> WebSocket-сервер. Клиент не работает с базой напрямую: все операции идут через серверный API, где выполняется валидация и контроль доступа.

## Главные компоненты

1. **Godot-клиент** — интерфейс входа, лобби, каталог, профиль, аватар.
2. **CloudAPI.gd** — HTTP-клиент внутри Godot. Через него идут login, profile, maps, catalog, friends, uploads.
3. **NetworkManager / client_network.gd** — подключение к игровому WebSocket-серверу.
4. **Node.js/Express API** — центральный серверный слой, который принимает запросы клиента.
5. **PocketBase** — база данных: профили, карты, предметы, друзья, инвентарь.
6. **Storage** — файлы: картинки, превью, текстуры, ассеты.
7. **Nginx** — публичная маршрутизация `/api`, `/ws`, `/downloads`.
8. **PM2** — держит API запущенным на VPS.

## Как работает вход

1. Игрок вводит логин и пароль.
2. Godot отправляет запрос на `/api/auth/v1/token`.
3. Node.js API проверяет пользователя в PocketBase.
4. Сервер возвращает access token.
5. Godot сохраняет сессию через UserSession.
6. Загружается профиль игрока.
7. Тяжёлая одежда и дополнительные avatar outfit данные догружаются фоном, чтобы лобби не зависало.

Что важно сказать: **логин не должен тянуть все данные мира сразу**. Мы специально разделили быстрый вход и фоновую загрузку тяжёлых данных.

## Как работает каталог карт

1. Пользователь создаёт или публикует карту.
2. Клиент отправляет данные на API.
3. API сохраняет запись в `maps`.
4. Превью карты хранится как файл или URL.
5. Лобби запрашивает список опубликованных карт.
6. Клиент показывает карточки игр.

Коллекция `maps` хранит не активные сервера, а постоянные данные карт. Активность игроков хранится отдельно в `active_servers`.

## Как работает активный сервер

1. Игровой сервер запускает комнату.
2. Он отправляет heartbeat в API.
3. API записывает состояние в `active_servers`.
4. Лобби видит, что в этой карте есть активная комната.
5. Игрок подключается к WebSocket-серверу.

## Как работает аватар и одежда

Есть два слоя данных:

- `profiles.avatar_data` — базовые данные профиля и аватара;
- `avatar_outfits` — актуальная одежда, цвета тела, надетые предметы.

Так сделано, чтобы профиль не становился огромным JSON. После добавления одежды раньше были подвисания, поэтому была сделана оптимизация: тяжёлые outfit-данные догружаются после входа фоном, а публичные списки друзей получают облегчённые профили.

## Как работает социальный модуль

Социальный модуль использует коллекции:

- `friend_requests` — заявки в друзья;
- `friendships` — подтверждённые дружбы;
- `follows` — подписки.

Клиент не может сам писать в эти таблицы напрямую. Он вызывает RPC-маршруты API, например отправку заявки или принятие заявки. Сервер сам определяет текущего пользователя по токену.

## Фактическая статистика

- auth-пользователи: {stats_lines['auth_users']}
- профили игроков: {stats_lines['profiles']}
- карты: {stats_lines['maps']}
- создатели карт: {stats_lines['map_creators']}
- посещения карт: {stats_lines['total_visits']}
- лайки карт: {stats_lines['total_likes']}
- модели: {stats_lines['model_assets']}
- avatar items: {stats_lines['avatar_items']}
- записи инвентаря: {stats_lines['inventory']}
- дружеские связи: {stats_lines['friendships']}
- размер PocketBase data: {stats_lines['pb_mb']} MB
- размер публичного storage: {stats_lines['storage_mb']} MB

## Что сказать про безопасность

- Клиент не подключается к базе напрямую.
- Все запросы проходят через Node.js API.
- Админ-статистика защищена токеном superuser.
- Service key не хранится в клиенте.
- Файлы сохраняются в контролируемой директории.
- Сервер обрезает слишком тяжёлые inline-картинки, чтобы клиент не зависал.
- Пользователь определяется по bearer token.

## Что сказать про нейросетевых агентов

Лучше формулировать так:

«При разработке использовались нейросетевые агенты как вспомогательный инструмент: для анализа ошибок, поиска узких мест и ускорения написания отдельных участков кода. Архитектура, постановка задач, проверка результата и интеграция выполнялись мной».

Не говори: «проект написала нейросеть». Говори: «нейросетевые агенты использовались как инструмент разработки».

## Возможные вопросы преподавателя

**Почему клиент не обращается к базе напрямую?**  
Чтобы не раскрывать доступы к базе, централизовать валидацию и не позволить клиенту напрямую менять чужие данные.

**Почему выбран Node.js?**  
Он хорошо подходит для сетевых API, быстро обрабатывает HTTP-запросы, имеет зрелую экосистему и удобно интегрируется с PocketBase.

**Зачем отдельно HTTP и WebSocket?**  
HTTP удобен для профилей, каталога и файлов, а WebSocket нужен для игровых событий в реальном времени.

**Что было самым сложным?**  
Синхронизация большого количества пользовательских данных без зависаний клиента: профиль, одежда, каталог, друзья, картинки.

**Как доказать, что база реально используется?**  
Показать статистику `/api/admin/stats`: сотни профилей, десятки карт, инвентарь, дружбы, файлы в storage.

## Демонстрация на защите

1. Показать вход в игру.
2. Показать профиль/лобби.
3. Показать каталог карт или предметов.
4. Показать, что данные сохраняются после перезахода.
5. Показать статистику сервера из JSON или админ-эндпоинта.
6. Коротко открыть схему архитектуры из отчёта.
"""
    GUIDE_MD.write_text(text, encoding="utf-8")


def main() -> None:
    stats = load_stats()
    stats_lines = stat_summary(stats)
    build_architecture_diagram()
    write_report_markdown(stats_lines)
    build_report_docx(stats_lines)
    write_defense_guide(stats_lines)
    print(REPORT_DOCX)
    print(REPORT_MD)
    print(GUIDE_MD)
    print(DIAGRAM_PATH)


if __name__ == "__main__":
    main()
